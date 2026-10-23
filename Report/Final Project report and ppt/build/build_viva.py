"""Viva-ready sheet: every component of the Dashboard, Signal Analysis and Diagnostics Studio tabs,
what it is, where its input comes from and how it is calculated.
Run: python3 build_viva.py  ->  ../SonicVitals_Viva_Sheet.docx (+ .pdf through Pages)"""
import os
import subprocess

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
SHOTS = os.path.join(HERE, 'viva_img')
# screenshot shown under the component's name (crops of the dashboard screenshots, made by hand from Images/)
SHOT = {
    'Blood-flow (perfusion) map': 'map.png', 'LIVENESS card': 'card_live.png', 'PERFUSION INDEX card': 'card_pi.png',
    'BREATHING PATTERN card': 'card_bpat.png', 'BLINK RATE card': 'ear.png', 'HEAD STEADINESS card': 'card_head.png',
    'Live EAR and breath plots': 'ear_breath.png', 'Recovered pulse plot': 'pulse.png',
    'Respiratory waveform plot': 'resp.png', 'Heart-rate history': 'history.png', 'Signal quality vs gate': 'quality.png',
    'Method comparison table and chart': 'methods.png',
    'rPPG power spectral density': 'psd.png', 'Chest microphone spectrogram': 'spectrogram.png',
    'Standardized pulse signal': 'stdpulse.png', 'Moving-average decomposition': 'decomp.png',
    'Skin colour change (R, G, B)': 'rgb.png', 'Method agreement over time': 'agree.png',
    'Snapshot frame': 'snapshot.png', 'Regions of interest': 'roi_list.png', 'Erythema index EI': 'idx_ei.png',
    'Pallor index PI': 'idx_pi_idx.png', 'Eye-band yellowness YI': 'idx_yi.png', 'Lip blue/red ratio CR': 'idx_cr.png',
    'Hemifacial asymmetry HFI': 'row_hfi.png', 'Frame clarity': 'row_clarity.png', 'Colour maps': 'colour_maps.png',
    'Index summary and observation cards': 'summary.png',
}
SHOT_W_CM, SHOT_H_CM = 4.4, 2.1
OUT = os.path.join(os.path.dirname(HERE), 'SonicVitals_Viva_Sheet.docx')
GREEN = RGBColor(0x0F, 0x6A, 0x4A)
HEAD_FILL = 'DCEBE3'

# (component, what it is, how it is obtained, how it is calculated / rule shown)
DASHBOARD = [
    ('Primary vision stream',
     'Live camera image with the face box and four skin boxes (forehead, both cheeks, nose); badge FULL FACE / '
     'alignment hint / SEARCHING.',
     'Mac webcam (30 fps, WebcamReader thread) or ESP32-S3 GC2145 camera (240 x 160, about 10 fps) over Wi-Fi.',
     'Haar-cascade face detection every 3rd frame; full-face gate (size, centring, both eyes). Between detections '
     'the boxes follow the face by sub-pixel phase correlation on a 96 x 96 Hann-windowed template (response >= 0.15, '
     'jump <= 0.12 face widths, relock after 8 misses or > 25 % drift).'),
    ('HEART RATE card',
     'One trustworthy heart rate in BPM with a status badge.',
     'Every method in the Method Comparison table (CHROM, POS, GREEN, Bayes tracker, TS-CAN, PhysNet, '
     'EfficientPhys) on the same face at the same moment.',
     'Robust consensus: median of all methods, keep those within 8 BPM, re-centre, average. Needs >= 3 methods and '
     '>= half agreeing for any value; green (confirmed) needs >= 4 methods, >= 60 % and an independent family (deep '
     'model) or the quality gate. Display stabiliser: EMA 0.3, jumps > 8 BPM held and flagged RAPID CHANGE until '
     'confirmed 4 times, hold 15 s, badge dwell 4 s.'),
    ('BREATHING RATE card',
     'Breaths per minute with badge.',
     'Four respiratory signals from the same video: skin brightness (RIIV), pulse height (RIAV), pulse rate (RIFV) '
     'and vertical head motion.',
     'Each source detrended (0.07 Hz), differenced (whitening), periodogram searched in 0.1-0.6 Hz (6-36 /min), edge '
     'peaks rejected, quality q per source (RIIV 0.45, motion 0.42, RIAV/RIFV 0.60). Reported when >= 2 sources agree '
     'within 4 /min (or one source with q >= 0.7), on >= 20 s of still data after the last head movement.'),
    ('HRV (RMSSD) card',
     'Short-term heart-rate variability, ms.',
     'Beat times of the camera pulse (or heart-sound intervals when the microphone is on the chest).',
     'CHROM pulse resampled to 30 Hz, band 0.7-3.5 Hz, spline x 8, peaks >= 0.6 beat apart; intervals within +-20 % '
     'of the median kept; RMSSD = sqrt(mean(diff(IBI)^2)). Shown only with >= 20 fps, >= 15 intervals, >= 80 % clean, '
     'pulse SNR >= +1 dB, result 5-150 ms, and its median HR within 8 BPM of the HR card.'),
    ('SIGNAL QUALITY card',
     'How dominant the heartbeat peak is in the pulse spectrum, %.',
     'Welch spectrum of the CHROM pulse (last 30 s).',
     'Q = power within +-0.1 Hz of f0 / power in 0.75-2.5 Hz x 100. Badge GOOD when the estimate passed the gate, '
     'FAIR when SNR is within 3 dB of the gate, POOR below that.'),
    ('HR SOURCE card',
     'How many methods the HR card averages.',
     'Output of the robust consensus.',
     'Count of methods inside the +-8 BPM band; badge ALL INCLUDED / N LEFT OUT (outliers named in the tooltip) / '
     'SINGLE METHOD.'),
    ('Chest microphone waveform',
     'Last 4 s of heart-sound (PCG) audio.',
     'INMP441 MEMS microphone, I2S at 16 kHz on the ESP32-S3, streamed over Wi-Fi.',
     '25-400 Hz Butterworth band-pass, decimated, min/max reduced to 600 points, normalised to +-1. Acoustic HR from the '
     'autocorrelation of the Shannon-energy envelope; used only when the rhythm check passes (rho >= 0.4, beat '
     'consistency >= 0.6). Never part of the HR consensus.'),
    ('Blood-flow (perfusion) map',
     'Colour map of where on the face the skin pulses with the heartbeat.',
     '30 s still capture started by the button; restarts if the face moves > 3 % or is lost > 3 s.',
     '32 x 32 skin cells; each cell\'s colour trace is locked-in at the card heart rate and projected onto the phase '
     'of the face-wide pulse (in-phase amplitude), Gaussian smoothed (sigma 1.5). Score z = (a - mean_off)/sd_off '
     'against off-frequency noise; cells with z >= 2.5 coloured, opacity by strength. Falls back to 16 or 8 cells '
     'when coverage < 30 %.'),
    ('LIVENESS card',
     'Whether a living face (not a photo or screen) is in front of the camera.',
     'Perfusion lock-in and blink detector.',
     'Evidence = share of skin cells pulsing x phase coherence. LIVE when evidence >= 0.15 on 2 updates or >= 2 '
     'blinks in 30 s; NO PULSE OR BLINK: PHOTO? when evidence < 0.08 on 3 updates with no blinks; otherwise CHECKING.'),
    ('PERFUSION INDEX card',
     'Strength of the skin pulse, %.',
     'Captured perfusion map.',
     'PI = 2 x A x 100 %, where A is the median in-phase amplitude (relative colour change) of the mapped cells, '
     'i.e. the peak-to-peak pulsatile fraction AC/DC.'),
    ('BREATHING PATTERN card',
     'REGULAR, IRREGULAR or PAUSE.',
     'Fused breathing waveform.',
     'Breath-to-breath intervals: coefficient of variation > 0.3 = IRREGULAR; envelope below 0.3 of its median for '
     '>= 8 s = PAUSE; otherwise REGULAR.'),
    ('BLINK RATE card',
     'Blinks per minute.',
     'MediaPipe Face Mesh (478 landmarks) on its own thread (newest frame), >= 15 fps.',
     'Eye aspect ratio EAR = (|p2-p6| + |p3-p5|) / (2|p1-p4|). A blink is a dip with prominence >= max(0.03, 0.35 x '
     'median EAR), width <= 0.5 s, >= 0.15 s after the previous one; rate = blinks in the window x 60 / window.'),
    ('HEAD STEADINESS card',
     'STEADY / SLIGHT MOVEMENT / MOVING.',
     'Motion-compensation shifts of the face (and head pose from solvePnP on mesh landmarks).',
     'RMS of face displacement over 10 s as % of face width: < 1.5 % STEADY, > 4 % MOVING, between = slight.'),
    ('Live EAR and breath plots',
     'Eye-aspect-ratio trace (last 20 s, blinks marked) and breathing waveform (last 30 s, breaths marked).',
     'Same sources as the blink and breathing cards.',
     'Plotted directly from the buffers; markers are the detected dips / breath peaks.'),
    ('Recovered pulse plot',
     'Blood-volume pulse recovered from skin colour.',
     'Skin-box RGB means at 15 Hz.',
     'CHROM: X = 3R - 2G, Y = 1.5R + G - 1.5B, S = X - (sigma_X / sigma_Y) Y; band-pass 0.75-2.5 Hz (3rd-order '
     'Butterworth, zero phase); z-scored.'),
    ('Respiratory waveform plot',
     'Slow breathing oscillation.',
     'Fused respiratory signal.',
     'Band-passed 0.1-0.6 Hz and z-scored.'),
    ('Heart-rate history',
     'Last 5 min of the HR card value with flags, each accepted (filled) and rejected (hollow) estimate, and the '
     'tracked HR.',
     'Stabiliser output, 2 s estimates and Bayes tracker.',
     'Tracker: HMM-style Bayes filter on a 0.5 BPM grid; posterior mass within +-3 BPM = confidence (>= 0.8 high).'),
    ('Signal quality vs gate',
     'Pulse SNR against the -2 dB acceptance line, plus how many of POS and GREEN agree.',
     'Welch spectrum of each 2 s estimate.',
     'SNR (de Haan) = 10 log10 [P(f0 +-0.1 Hz) + P(2f0 +-0.2 Hz)] / [rest of 0.5-4 Hz]. Accepted when SNR >= -2 dB, '
     'POS and GREEN within 6 BPM (2/2), >= 80 % real frames, <= 15 % repaired.'),
    ('Method comparison table and chart',
     'Heart rate from every method at the same moment, marked IN AVERAGE / OUTLIER / NOT USED FOR HR.',
     'CHROM, POS, GREEN every 2 s; deep models on the last 30 s of 72 x 72 face crops every 12 s.',
     'Same consensus as the HR card; dotted line = median, spread shown. Deep models: TS-CAN, PhysNet, EfficientPhys '
     '(rPPG-Toolbox, trained on PURE), run on CPU.'),
]

ANALYSIS = [
    ('CARDIAC PEAK (f0)',
     'Frequency of the heartbeat peak.',
     'CHROM spectrum of the last 30 s of gated frames.',
     'Highest peak in 0.75-2.5 Hz after peak tracking and lighting-peak rejection; BPM = 60 x f0. Needs 10 s of data.'),
    ('SNR (DE HAAN)',
     'Signal-to-noise ratio of the pulse, dB.',
     'Same spectrum.',
     '10 log10 of power at f0 (+-0.1 Hz) and 2f0 (+-0.2 Hz) over the rest of 0.5-4 Hz; gate -2 dB.'),
    ('SPECTRAL CONCENTRATION (Q)',
     'Share of cardiac-band power at the peak, %.',
     'Same spectrum.',
     'P(f0 +-0.1 Hz) / P(0.75-2.5 Hz) x 100.'),
    ('RESPIRATION (fr)',
     'Breathing frequency, Hz and breaths/min.',
     'Four-source breathing fusion.',
     'fr = RR / 60.'),
    ('METHODS AGREEING | MOTION',
     'Consensus count, repaired share and skipped time.',
     'Quality gate of the 2 s estimate.',
     'POS and GREEN within 6 BPM of CHROM (0-2); % of frame pairs repaired by the derivative-domain artifact filter; '
     'seconds skipped before the last still stretch.'),
    ('VIDEO SOURCE (MEASURED FPS)',
     'Which camera and the processed frame rate.',
     'Frame timestamps.',
     'Frames processed per second over the last window.'),
    ('rPPG power spectral density',
     'Pulse power against frequency with the cardiac band shaded and the chosen peak marked.',
     'High-passed (0.6 Hz) CHROM pulse at 15 Hz.',
     'Welch periodogram, 256-sample segments, NFFT 4096, dB relative to the maximum, shown 0.5-3.5 Hz.'),
    ('Chest microphone spectrogram',
     'Time-frequency picture of the last 4 s of chest audio.',
     'INMP441 audio.',
     'STFT, Hamming window about 32 ms, 75 % overlap, up to 2 kHz, 60 time columns, colour -60 to 0 dB. Heart sounds '
     'appear as short vertical stripes below about 200 Hz at the heart rate.'),
    ('Standardized pulse signal',
     'The recovered pulse, enlarged.',
     'CHROM pulse.',
     'Band-pass 0.75-2.5 Hz, z-score (mean 0, SD 1).'),
    ('Moving-average decomposition',
     'Trend, pulsatile part and first difference of the pulse.',
     'CHROM pulse.',
     'Trend = moving average over W = min(14, N/4) samples (about 1 s); pulsatile = x - trend; first difference '
     'x[n] - x[n-1]; pulsatile/trend energy = motion ratio.'),
    ('Skin colour change (R, G, B)',
     'Percentage change of the skin-box mean colours over 20 s: the raw input of the chain.',
     'Skin-masked, motion-compensated box means after artifact repair and illumination rectification.',
     '(channel / its mean - 1) x 100. The pulse is about 0.1-0.5 % and strongest in green.'),
    ('Method agreement over time',
     'HR from CHROM, POS and GREEN at every update over 5 min.',
     'Each extractor through the same spectrum and peak search.',
     'POS: S1 = G - B, S2 = -2R + G + B, P = S1 + (sigma_S1 / sigma_S2) S2 (1.6 s windows). GREEN = detrended green. '
     'Overlapping lines = real pulse.'),
    ('EXPORT RECORDING (CSV)',
     'Download of the display log.',
     'Browser display log, one row per update.',
     'Columns: time, HR card value and state, RR card value and state, accepted HR/RR, RMSSD, HR source, optical and '
     'acoustic HR, quality, video source.'),
]

DIAGNOSTICS = [
    ('Snapshot frame',
     'The image the indices are computed on.',
     'Sharpest full-face frame of the last 8 s (or the live frame); AUTO-CAPTURE triggers once a full face is locked.',
     'Sharpness = variance of the Laplacian of the grey image; the frame with the highest value in the window is kept.'),
    ('Pipeline strip (5 steps)',
     'Frame Captured -> ROI Segmentation -> ROI Colour Means -> Index Calculation -> Range Check.',
     'Progress of the snapshot analysis.',
     'Each step lights when its stage completes.'),
    ('Regions of interest',
     'Five regions relative to the face box (x, y, w, h as fractions of the face).',
     'Face box of the snapshot.',
     'Forehead (0.20, 0.08, 0.60, 0.18); left cheek (0.12, 0.44, 0.26, 0.24); right cheek (0.62, 0.44, 0.26, 0.24); '
     'eye band (0.20, 0.28, 0.60, 0.14); lips (0.28, 0.74, 0.44, 0.18). Mean R, G, B per region.'),
    ('Erythema index EI',
     'Redness of the forehead skin.',
     'Forehead mean R and G.',
     'EI = 100 (log10 R - log10 G). Software range 8.0-23.0.'),
    ('Pallor index PI',
     'Paleness of the cheeks.',
     'Both cheeks\' means.',
     'PI = (G_L + B_L + G_R + B_R) / (2 (R_L + R_R)). In range < 0.72, borderline <= 0.84.'),
    ('Eye-band yellowness YI',
     'Yellowness around the eyes (band includes eyelid skin, not only sclera).',
     'Eye-band means.',
     'YI = 100 (R + G - 2B) / (R + G + B). In range < 12, borderline <= 18.'),
    ('Lip blue/red ratio CR',
     'Blueness of the lips relative to red.',
     'Lip means.',
     'CR = B / R. In range < 0.76, borderline <= 0.86.'),
    ('Hemifacial asymmetry HFI',
     'Left/right colour difference (also reveals uneven lighting).',
     'Cheek green means.',
     'HFI = 100 |G_L - G_R| / (G_L + G_R). In range < 5 %, borderline <= 10 %.'),
    ('Frame clarity',
     'Focus / motion-blur check of the snapshot.',
     'Snapshot grey image.',
     'Laplacian variance; > 100 = ADEQUATE FOCUS, else BLUR / MOTION.'),
    ('Colour maps',
     'Per-pixel views of the face: R - G (jet), CIELAB b* (inferno), B / R (cool), plus ROI and raw views.',
     'Face crop of the snapshot.',
     'Value per pixel, min-max normalised within the face, colour-mapped and blended 55 % over the image. Shows '
     'relative differences inside one image only.'),
    ('Index summary and observation cards',
     'Table of the five indices with value, software range and status; flags count.',
     'Index calculation above.',
     'Each index classified IN RANGE / BORDERLINE / OUT OF RANGE; "N OF 5 OUT OF RANGE". Uncalibrated 8-bit camera '
     'values: exploratory screening aids, not diagnoses.'),
]

QA = [
    ('Why does a camera see the heartbeat?',
     'Each beat raises blood volume in skin capillaries; haemoglobin absorbs green light most, so skin reflects '
     'about 0.1-0.5 % less light at systole. Averaging many skin pixels makes this visible above camera noise.'),
    ('Why CHROM / POS instead of just green?',
     'Motion and lighting change all three channels together (specular / intensity), while the pulse has a specific '
     'colour direction. CHROM and POS project RGB onto a plane that cancels intensity changes, so they reject motion '
     'far better than GREEN (UBFC MAE 2.40 / 2.08 vs 22.04 BPM).'),
    ('What does the quality gate prevent?',
     'Showing a number when the spectrum has no dominant peak. It withholds the 19 % hardest UBFC windows and lowers '
     'MAE from 2.30 to 1.56 BPM.'),
    ('Why average several methods for the card?',
     'Independent methods only agree by chance rarely; the median rejects outliers, and requiring a deep model '
     'family for a confirmed value avoids all classic methods agreeing on the same lighting artifact.'),
    ('How is breathing obtained without a chest sensor?',
     'Breathing modulates skin brightness, pulse height, pulse rate and moves the head slightly. Four signals are '
     'searched independently and the rate is reported only when at least two agree.'),
    ('Why is camera HRV restricted?',
     'RMSSD depends on beat timing to a few ms; at low SNR camera noise shifts the peaks (error 64 ms at -2 dB vs '
     '11 ms at +1 dB), so HRV is shown only at SNR >= +1 dB and >= 20 fps.'),
    ('What is skin-metered exposure control?',
     'The ESP32 sets the GC2145 exposure from the brightness of the skin pixels, not the whole frame, so the face '
     'stays in the sensor\'s linear range and is neither clipped nor too dark.'),
    ('How does liveness reject a photo?',
     'A printed face has no in-phase pulsation across the skin and does not blink; photos gave 13/13 NO PULSE and '
     'pulsing faces 5/5 LIVE.'),
    ('Are the Diagnostics Studio indices medical?',
     'No. They are uncalibrated colour ratios with software ranges; they need calibration and clinical validation.'),
]


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    el.set(qn('w:val'), 'true')
    trPr.append(el)


def section_table(doc, title, rows):
    # the section title is the table's first (merged) row: Pages starts every table on a new page,
    # so a separate heading paragraph would be left alone on the page before
    widths = [Cm(5.0), Cm(5.0), Cm(5.5), Cm(11.6)]
    t = doc.add_table(rows=2, cols=4)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    top = t.rows[0].cells[0].merge(t.rows[0].cells[3])
    top.text = ''
    r = top.paragraphs[0].add_run(title)
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    shade(top, '0F6A4A')
    for i, h in enumerate(['Component', 'What it is', 'How it is obtained (input)', 'How it is calculated']):
        c = t.rows[1].cells[i]
        c.text = ''
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        shade(c, HEAD_FILL)
    repeat_header(t.rows[0])
    repeat_header(t.rows[1])
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ''
            r = cells[i].paragraphs[0].add_run(val)
            if i == 0:
                r.bold = True
                if row[0] in SHOT:
                    f = os.path.join(SHOTS, SHOT[row[0]])
                    with Image.open(f) as im:
                        aspect = im.height / im.width
                    w = min(SHOT_W_CM, SHOT_H_CM / aspect)
                    cells[0].add_paragraph().add_run().add_picture(f, width=Cm(w))
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = w
    # Pages lays tables out from the grid, not the cell widths: write the grid and a fixed layout too
    t.autofit = False
    for gc, w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), widths):
        gc.set(qn('w:w'), str(int(w.twips)))
    doc.add_paragraph()  # keeps consecutive tables separate


def heading(doc, text, level=1, new_page=False):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.page_break_before = new_page
    for r in p.runs:
        r.font.color.rgb = GREEN


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for m in ('left_margin', 'right_margin'):
        setattr(sec, m, Cm(1.3))
    sec.top_margin = sec.bottom_margin = Cm(1.3)
    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(9.5)
    st.paragraph_format.space_after = Pt(2)

    title = doc.add_heading('SonicVitals: viva sheet', 0)
    for r in title.runs:
        r.font.color.rgb = GREEN
    doc.add_paragraph('Every component of the Dashboard, Signal Analysis and Diagnostics Studio tabs: what it is, '
                      'what it is computed from, and how. Janhavee Dhote (24BML0120), Pratik Kumar (24BML0027), '
                      'Uday Krishan Khandelwal (24BML0012). Guide: Jeeva J. B.')

    heading(doc, 'Key numbers to remember', 1)
    for line in [
        'Pulse band 0.75-2.5 Hz (45-150 BPM); SNR gate -2 dB; POS and GREEN within 6 BPM; >= 80 % real frames; <= 15 % repaired.',
        'Consensus: +-8 BPM band; >= 3 methods for a value; >= 4 methods, 60 % and an independent family for CONFIRMED.',
        'Stabiliser: jump 8 BPM, EMA 0.3, confirm 4, hold 15 s, badge dwell 4 s.',
        'Breathing: 6-36 /min, 4 sources, agree within 4 /min, >= 20 s still data.',
        'HRV: >= 20 fps, SNR >= +1 dB, >= 15 intervals, >= 80 % clean, 5-150 ms.',
        'UBFC-rPPG: MAE 1.56 BPM, RMSE 2.73, r = 0.99, 93 % within 5 BPM (gated).',
        'Live: 27 fps processed, 10.6 ms per frame; perfusion capture 30 s, z >= 2.5.',
    ]:
        doc.add_paragraph(line, style='List Bullet')

    section_table(doc, '1. Dashboard', DASHBOARD)
    section_table(doc, '2. Signal Analysis', ANALYSIS)
    section_table(doc, '3. Diagnostics Studio', DIAGNOSTICS)

    doc.add_paragraph()
    heading(doc, '4. Likely viva questions', 1)
    for q, a in QA:
        p = doc.add_paragraph()
        p.add_run(q + ' ').bold = True
        p.add_run(a)

    doc.save(OUT)
    print(OUT)
    pdf = OUT[:-5] + '.pdf'
    subprocess.run([os.path.join(HERE, 'pages_pdf.sh'), OUT, pdf], check=False)
    print(pdf if os.path.exists(pdf) else 'PDF not produced (needs Pages running)')


if __name__ == '__main__':
    main()
